#!/usr/bin/env python3
"""Offline localization validator / CI gate.

Validates existing translations against the source copy WITHOUT calling any model
or network, so it is fast, deterministic, and free — ideal for a PR check.

Blocking issues (non-zero exit):
- Placeholder/token mismatch between source and a non-empty translation.
- Missing translations, only when --require-complete is passed.

Advisory issues (reported, never fail): glossary term not applied, extreme
length ratio, locale keys absent from the source.

Usage:
  python3 loc_validate.py --marketing marketing_config.json
  python3 loc_validate.py --workbook config.json --require-complete
  python3 loc_validate.py --marketing marketing_config.json --report report.json
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

import loc_core as core


def _add(report, level, **fields):
    report[level].append(fields)


def validate_marketing(config_path: Path, require_complete: bool, report: dict) -> None:
    cfg = core.load_json(config_path)
    base = config_path.parent
    outdir = base / cfg["output_dir"]
    depth = int(cfg.get("section_depth", 1))

    # Prefer corrected English if the pipeline has produced it.
    corrected_path = outdir / cfg.get("corrected_english_file", "en.corrected.json")
    source_path = corrected_path if corrected_path.exists() else base / cfg["source_file"]
    source = core.flatten(core.load_json(source_path, {}))

    core_glossary = core.load_json(base / cfg["core_glossary_file"], {})
    marketing_glossary = core.load_json(base / cfg["marketing_glossary_file"], {})

    for lang_code, filename in cfg["locales"].items():
        lang = core.canonical_language(lang_code)
        target_path = outdir / filename
        target = core.flatten(core.load_json(target_path, {}))
        glossary = core.glossary_for(lang, core_glossary, marketing_glossary)

        for key, src in source.items():
            if not isinstance(src, str) or not src.strip():
                continue
            val = target.get(key)
            if val in (None, ""):
                if require_complete:
                    _add(report, "blocking", kind="missing", scope=filename, key=key)
                else:
                    _add(report, "advisory", kind="missing", scope=filename, key=key)
                continue
            if not core.tokens_match(src, val):
                _add(report, "blocking", kind="placeholder", scope=filename, key=key,
                     source_tokens=core.tokens(src), target_tokens=core.tokens(val))
                continue
            for v in core.glossary_violations(src, val, glossary):
                _add(report, "advisory", kind="glossary", scope=filename, key=key, detail=v)
            ratio = core.length_ratio_flag(src, val)
            if ratio:
                _add(report, "advisory", kind="length", scope=filename, key=key, detail=ratio)

        for key in target:
            if key not in source:
                _add(report, "advisory", kind="orphan", scope=filename, key=key)


def validate_workbook(config_path: Path, require_complete: bool, report: dict) -> None:
    from openpyxl import load_workbook  # lazy: only needed for workbook mode

    cfg = core.load_json(config_path)
    base = config_path.parent
    wb_path = base / cfg.get("output_file", "UI_translations_localized.xlsx")
    if not wb_path.exists():
        wb_path = base / cfg["input_file"]
    worksheet = cfg.get("worksheet", "translations")
    wb = load_workbook(wb_path, read_only=True)
    ws = wb[worksheet]

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    lut = {str(h).strip().lower(): i for i, h in enumerate(headers) if h is not None}

    def col(name, *alts):
        for n in (name, *alts):
            if n and n.strip().lower() in lut:
                return lut[n.strip().lower()]
        return None

    key_i = col(cfg.get("key_column", "key"))
    en_i = col(cfg.get("source_column", "en"), "English")
    corr_i = col(cfg.get("english_review", {}).get("corrected_column", "corrected_en"))
    glossary = core.load_json(base / cfg["glossary_file"], {}) if cfg.get("glossary_file") else {}
    targets = {}
    for lang in ("ja", "tw", "kr"):
        configured = cfg.get("target_columns", {}).get(lang)
        idx = col(configured or "", *{"ja": ["jp", "ja"], "tw": ["tw", "zh-TW"], "kr": ["kr", "ko"]}[lang])
        if idx is not None:
            targets[lang] = idx

    for r in ws.iter_rows(min_row=2):
        src = (corr_i is not None and r[corr_i].value) or (en_i is not None and r[en_i].value)
        if not isinstance(src, str) or not src.strip():
            continue
        key = str(r[key_i].value) if key_i is not None and r[key_i].value is not None else ""
        for lang, idx in targets.items():
            val = r[idx].value
            if val in (None, ""):
                lvl = "blocking" if require_complete else "advisory"
                _add(report, lvl, kind="missing", scope=lang, key=key)
                continue
            if not core.tokens_match(src, str(val)):
                _add(report, "blocking", kind="placeholder", scope=lang, key=key,
                     source_tokens=core.tokens(src), target_tokens=core.tokens(str(val)))
                continue
            terms = core.glossary_for(lang, glossary)
            for v in core.glossary_violations(src, str(val), terms):
                _add(report, "advisory", kind="glossary", scope=lang, key=key, detail=v)
            ratio = core.length_ratio_flag(src, str(val))
            if ratio:
                _add(report, "advisory", kind="length", scope=lang, key=key, detail=ratio)
    wb.close()


def main() -> int:
    ap = argparse.ArgumentParser(description="Offline localization CI gate.")
    ap.add_argument("--marketing", help="Path to a marketing_config.json")
    ap.add_argument("--workbook", help="Path to a workbook config.json")
    ap.add_argument("--require-complete", action="store_true",
                    help="Treat missing translations as blocking.")
    ap.add_argument("--report", help="Write the full JSON report to this path.")
    args = ap.parse_args()

    if not args.marketing and not args.workbook:
        ap.error("Provide --marketing and/or --workbook")

    report: dict[str, list] = {"blocking": [], "advisory": []}
    if args.marketing:
        validate_marketing(Path(args.marketing), args.require_complete, report)
    if args.workbook:
        validate_workbook(Path(args.workbook), args.require_complete, report)

    if args.report:
        core.save_json_atomic(Path(args.report), report)

    by_kind: dict[str, int] = {}
    for item in report["blocking"] + report["advisory"]:
        by_kind[item["kind"]] = by_kind.get(item["kind"], 0) + 1

    print("Localization validation")
    print(f"  blocking: {len(report['blocking'])}  advisory: {len(report['advisory'])}")
    for kind, n in sorted(by_kind.items()):
        print(f"    {kind}: {n}")
    for item in report["blocking"][:50]:
        print(f"  BLOCK [{item['kind']}] {item.get('scope')} :: {item.get('key')} "
              f"{item.get('source_tokens', '')}{(' -> ' + str(item['target_tokens'])) if 'target_tokens' in item else ''}")
    if len(report["blocking"]) > 50:
        print(f"  ... and {len(report['blocking']) - 50} more blocking issues")

    if report["blocking"]:
        print("FAILED: blocking localization issues found.")
        return 1
    print("OK: no blocking localization issues.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
