#!/usr/bin/env python3
"""KBKx workbook localization pipeline (v6).

Translates the platform/admin UI string workbook (``en`` -> ``ja``/``tw``/``kr``).

Pipeline:
1. Normalize English into ``corrected_en``.
2. Translate from ``corrected_en``.
3. Re-review existing translations when corrected English changes.
4. Save after each batch and resume safely.

v6 changes:
- Shared logic lives in ``loc_core``.
- Structured (JSON-schema) model outputs with graceful fallback.
- Concurrent batches (``max_workers``) with main-thread checkpointing.
- Advisory glossary-compliance and length-ratio QA flags.
"""
from __future__ import annotations

import argparse
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import loc_core as core

DEFAULT_TARGET_HEADERS = {
    "ja": ["jp", "ja", "Japanese"],
    "tw": ["tw", "zh-TW", "Traditional Chinese (Taiwan)"],
    "kr": ["kr", "ko", "Korean"],
}


@dataclass
class RowItem:
    row: int
    key: str
    original_english: str
    corrected_english: str = ""


def find_header(headers: list[Any], candidates: list[str]) -> int:
    lookup = {str(h).strip().lower(): i + 1 for i, h in enumerate(headers) if h is not None}
    for candidate in candidates:
        found = lookup.get(candidate.strip().lower())
        if found:
            return found
    raise ValueError(f"Missing columns {candidates}; found {headers}")


def find_or_create_header(ws, header: str) -> int:
    headers = [c.value for c in ws[1]]
    lookup = {str(h).strip().lower(): i + 1 for i, h in enumerate(headers) if h is not None}
    found = lookup.get(header.strip().lower())
    if found:
        return found
    col = ws.max_column + 1
    ws.cell(1, col).value = header
    return col


def english_instructions(glossary: dict[str, Any]) -> str:
    concepts = "\n".join(f"- {k}" for k in glossary.keys())
    return f"""You are the senior UX writer for KBKx, an enterprise B2B travel and hotel distribution platform.

Rewrite each English backend/admin UI string into concise, natural enterprise SaaS English.

Rules:
1. Return the result as items; each item has exactly row, corrected_english, changed, note.
2. Preserve meaning, business behavior, placeholders, ICU MessageFormat, HTML tags, punctuation tokens, % tokens and \\n.
3. Correct grammar, unnatural wording, inconsistent capitalization and non-native phrasing.
4. Do not simplify away meaningful distinctions such as Confirm, Submit, Complete, Distribution Channel, Distribution Partner and Distributor.
5. Standardize approved product concepts where the original wording clearly refers to the same concept.
6. Do not rewrite proper nouns, technical IDs, API names or product names.
7. changed must be true only when corrected_english differs from the original.
8. Use sentence case for messages and concise title case or sentence case appropriate to UI labels.

Approved KBKx concepts:
{concepts}
"""


def translation_instructions(language: str, glossary: dict[str, Any]) -> str:
    terms = core.glossary_for(language, glossary)
    glossary_text = "\n".join(f"- {k} => {v}" for k, v in terms.items())
    return f"""You are the localization lead for KBKx, an enterprise B2B travel and hotel distribution platform.

Translate corrected enterprise UI English into {core.LANGUAGE_NAMES[language]}.

Rules:
1. Return the result as items; each item has exactly row, translation, note.
2. Translate from corrected_english, not original_english.
3. Preserve every placeholder, ICU MessageFormat expression, HTML tag, % token and \\n exactly.
4. Use concise, natural hospitality and enterprise SaaS terminology.
5. Keep Distribution Channel, Distribution Partner and Distributor distinct.
6. Translate Source according to context: inventory source, pricing source, content source, supplier, property or internal/KAOS.
7. Japanese may use 予約 for both Reservation and Booking where natural.
8. Taiwan output must use Taiwan Traditional Chinese terminology, not Simplified Chinese.
9. Do not include explanations in translation; use note for ambiguity.

Approved glossary:
{glossary_text}
"""


def append_qa(qa, language, row, key, original, corrected, translation, issue):
    qa.append([language, row, key, original, corrected, translation, issue])


def grouped_representatives(items, text_selector, depth):
    groups: dict[str, list[RowItem]] = {}
    for item in items:
        groups.setdefault(core.cache_key(text_selector(item), item.key, depth), []).append(item)
    representatives = [members[0] for members in groups.values()]
    return representatives, groups


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--input")
    parser.add_argument("--output")
    parser.add_argument("--languages", nargs="+")
    parser.add_argument("--model")
    parser.add_argument("--batch-size", type=int)
    parser.add_argument("--max-workers", type=int)
    parser.add_argument("--batch", action="store_true", help="Use the async Batch API (cheaper, slower).")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    config_path = Path(args.config)
    config = core.load_json(config_path)
    base = config_path.parent
    core.load_env_file(base / ".env")  # optional; never overrides a set env var

    input_path = Path(args.input or config.get("input_file", "translations.xlsx"))
    output_path = Path(args.output or config.get("output_file", "translations_localized.xlsx"))
    if not input_path.is_absolute():
        input_path = base / input_path
    if not output_path.is_absolute():
        output_path = base / output_path

    if str(config.get("processing_mode", "standard")).lower() != "standard":
        raise ValueError("This implementation uses Standard Responses API processing. Set processing_mode to standard.")

    model = args.model or config.get("model", "gpt-5.4-mini")
    batch_size = args.batch_size or int(config.get("batch_size", 30))
    max_workers = args.max_workers or int(config.get("max_workers", 4))
    use_batch = args.batch or bool(config.get("batch_api", False))
    batch_opts = {
        "poll_interval": float(config.get("batch_poll_interval_seconds", 15)),
        "timeout": float(config.get("batch_timeout_seconds", 86400)),
    }
    languages = args.languages or config.get("languages", ["ja", "tw", "kr"])
    languages = [core.canonical_language(lang) for lang in languages]
    overwrite = args.overwrite or not bool(config.get("skip_existing", True))
    review_cfg = config.get("english_review", {})
    worksheet = config.get("worksheet", "translations")

    cache_cfg = config.get("cache", {})
    cache_enabled = bool(cache_cfg.get("enabled", True))
    context_depth = int(cache_cfg.get("context_depth", 2))
    cache_file = base / cache_cfg.get("file", "localization_cache.json")
    cache = core.load_json(cache_file, {"version": 1, "english": {}, "translations": {}})
    cache.setdefault("english", {})
    cache.setdefault("translations", {})
    for lang in ("ja", "tw", "kr"):
        cache["translations"].setdefault(lang, {})

    glossary_file = config.get("glossary_file")
    glossary = core.load_glossary(base / glossary_file if glossary_file else None)

    provider = core.build_provider(config, model_override=model)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path != input_path and not output_path.exists():
        shutil.copy2(input_path, output_path)

    workbook_path = output_path if output_path.exists() else input_path
    wb = load_workbook(workbook_path)
    if worksheet not in wb.sheetnames:
        raise ValueError(f"Worksheet {worksheet!r} not found")
    ws = wb[worksheet]

    headers = [c.value for c in ws[1]]
    key_col = find_header(headers, [config.get("key_column", "key")])
    en_col = find_header(headers, [config.get("source_column", "en"), "English"])

    corrected_col = find_or_create_header(ws, review_cfg.get("corrected_column", "corrected_en"))
    status_col = find_or_create_header(ws, review_cfg.get("english_status_column", "english_status"))
    fingerprint_col = find_or_create_header(ws, review_cfg.get("english_fingerprint_column", "english_fingerprint"))
    review_existing = bool(
        review_cfg.get("review_existing_translations_when_english_changes", True)
        or config.get("review_existing_translations_when_english_changes", True)
    )
    glossary_check = bool(config.get("qa", {}).get("check_glossary_compliance", True))
    length_check = bool(config.get("qa", {}).get("check_length_ratio", True))

    headers = [c.value for c in ws[1]]
    target_cols = {}
    for language in ("ja", "tw", "kr"):
        configured = config.get("target_columns", {}).get(language)
        target_cols[language] = find_header(
            headers,
            ([configured] if configured else []) + DEFAULT_TARGET_HEADERS[language],
        )

    qa_name = config.get("qa", {}).get("report_sheet_name", "Localization QA")
    if qa_name in wb.sheetnames:
        qa = wb[qa_name]
    else:
        qa = wb.create_sheet(qa_name)
        qa.append(["Language", "Row", "Key", "Original English", "Corrected English", "Translation", "Issue / Note"])

    stats = {
        "english_cache_hits": 0,
        "translation_cache_hits": {"ja": 0, "tw": 0, "kr": 0},
        "english_api_items": 0,
        "translation_api_items": {"ja": 0, "tw": 0, "kr": 0},
        "deduplicated_rows": 0,
    }

    def checkpoint():
        wb.save(output_path)
        if cache_enabled:
            core.save_json_atomic(cache_file, cache)

    # ----------------------------- PHASE 1: English ----------------------------- #
    english_pending = []
    for row in range(2, ws.max_row + 1):
        original = ws.cell(row, en_col).value
        if not isinstance(original, str) or not original.strip():
            continue
        old_fp = ws.cell(row, fingerprint_col).value
        current_fp = core.fingerprint(original)
        corrected = ws.cell(row, corrected_col).value
        if corrected in (None, "") or old_fp != current_fp:
            english_pending.append(RowItem(row, str(ws.cell(row, key_col).value or ""), original))

    changed_rows: set[int] = set()
    uncached_english = []
    for item in english_pending:
        ck = core.cache_key(item.original_english, item.key, context_depth)
        cached = cache["english"].get(ck) if cache_enabled else None
        if cached:
            corrected = str(cached["corrected_english"])
            previous = ws.cell(item.row, corrected_col).value
            ws.cell(item.row, corrected_col).value = corrected
            ws.cell(item.row, fingerprint_col).value = core.fingerprint(item.original_english)
            ws.cell(item.row, status_col).value = "changed" if corrected != item.original_english else "unchanged"
            if corrected != item.original_english or (previous not in (None, "") and previous != corrected):
                changed_rows.add(item.row)
            stats["english_cache_hits"] += 1
        else:
            uncached_english.append(item)

    representatives, english_groups = grouped_representatives(
        uncached_english, lambda x: x.original_english, context_depth
    )
    stats["deduplicated_rows"] += len(uncached_english) - len(representatives)

    print(
        f"English normalization: {len(english_pending)} pending; "
        f"{stats['english_cache_hits']} cache hits; {len(representatives)} unique API items "
        f"(workers={max_workers})",
        flush=True,
    )

    batches = core.chunk(representatives, batch_size)
    english_prompt = english_instructions(glossary)
    specs = [
        {"instructions": english_prompt,
         "payload": [{"row": x.row, "key": x.key, "english": x.original_english} for x in b],
         "schema": core.ENGLISH_ROW_SCHEMA, "schema_name": "english_normalization"}
        for b in batches
    ]

    total = len(batches)
    for i, results in core.dispatch(provider, specs, max_workers, use_batch, **batch_opts):
        batch = batches[i]
        print(f"  English batch {i + 1}/{total}: {len(batch)} unique strings", flush=True)
        by_row = {int(x["row"]): x for x in results if "row" in x}
        stats["english_api_items"] += len(batch)

        for rep in batch:
            result = by_row.get(rep.row)
            group_key = core.cache_key(rep.original_english, rep.key, context_depth)
            members = english_groups[group_key]
            if not result:
                for item in members:
                    append_qa(qa, "en", item.row, item.key, item.original_english, "", "", "Missing English result")
                continue

            corrected = str(result.get("corrected_english", rep.original_english)).strip() or rep.original_english
            note = str(result.get("note", "")).strip()
            if not core.tokens_match(rep.original_english, corrected):
                corrected = rep.original_english
                note = (note + "; " if note else "") + "English placeholder mismatch; original retained"

            cache["english"][group_key] = {"corrected_english": corrected, "note": note}

            for item in members:
                previous = ws.cell(item.row, corrected_col).value
                ws.cell(item.row, corrected_col).value = corrected
                ws.cell(item.row, fingerprint_col).value = core.fingerprint(item.original_english)
                changed = corrected != item.original_english
                ws.cell(item.row, status_col).value = "changed" if changed else "unchanged"
                if changed or (previous not in (None, "") and previous != corrected):
                    changed_rows.add(item.row)
                if changed or note:
                    append_qa(qa, "en", item.row, item.key, item.original_english, corrected, "", note)

        checkpoint()
        print(f"  Saved English checkpoint to {output_path.name}", flush=True)

    # --------------------------- PHASE 2: Translation --------------------------- #
    for lang in languages:
        target_col = target_cols[lang]
        translation_prompt = translation_instructions(lang, glossary)
        terms = core.glossary_for(lang, glossary) if glossary_check else {}
        pending = []

        for row in range(2, ws.max_row + 1):
            original = ws.cell(row, en_col).value
            corrected = ws.cell(row, corrected_col).value or original
            target = ws.cell(row, target_col).value
            key = str(ws.cell(row, key_col).value or "")
            if not isinstance(corrected, str) or not corrected.strip():
                continue
            needs_translation = overwrite or target in (None, "")
            needs_review = review_existing and row in changed_rows and target not in (None, "")
            if needs_translation or needs_review:
                pending.append(RowItem(row, key, str(original or ""), str(corrected)))

        uncached = []
        for item in pending:
            ck = core.cache_key(item.corrected_english, item.key, context_depth)
            cached = cache["translations"][lang].get(ck) if cache_enabled else None
            if cached:
                translation = str(cached["translation"])
                if core.tokens_match(item.corrected_english, translation):
                    existing = ws.cell(item.row, target_col).value
                    ws.cell(item.row, target_col).value = translation
                    if existing not in (None, "") and existing != translation:
                        append_qa(
                            qa, lang, item.row, item.key, item.original_english,
                            item.corrected_english, translation,
                            "Existing translation revised from translation memory",
                        )
                    stats["translation_cache_hits"][lang] += 1
                    continue
            uncached.append(item)

        representatives, groups = grouped_representatives(
            uncached, lambda x: x.corrected_english, context_depth
        )
        stats["deduplicated_rows"] += len(uncached) - len(representatives)

        print(
            f"{core.LANGUAGE_NAMES[lang]}: {len(pending)} pending/review; "
            f"{stats['translation_cache_hits'][lang]} cache hits; "
            f"{len(representatives)} unique API items (workers={max_workers})",
            flush=True,
        )

        batches = core.chunk(representatives, batch_size)
        specs = [
            {"instructions": translation_prompt,
             "payload": [
                 {
                     "row": x.row,
                     "key": x.key,
                     "original_english": x.original_english,
                     "corrected_english": x.corrected_english,
                     "existing_translation": ws.cell(x.row, target_col).value or "",
                     "review_existing": ws.cell(x.row, target_col).value not in (None, ""),
                 }
                 for x in b
             ],
             "schema": core.TRANSLATION_ROW_SCHEMA, "schema_name": "translation"}
            for b in batches
        ]

        total = len(batches)
        for i, results in core.dispatch(provider, specs, max_workers, use_batch, **batch_opts):
            batch = batches[i]
            print(f"  {core.LANGUAGE_NAMES[lang]} batch {i + 1}/{total}: {len(batch)} unique strings", flush=True)
            by_row = {int(x["row"]): x for x in results if "row" in x}
            stats["translation_api_items"][lang] += len(batch)

            for rep in batch:
                result = by_row.get(rep.row)
                group_key = core.cache_key(rep.corrected_english, rep.key, context_depth)
                members = groups[group_key]

                if not result:
                    for item in members:
                        append_qa(
                            qa, lang, item.row, item.key, item.original_english,
                            item.corrected_english, "", "Missing translation result",
                        )
                    continue

                translation = str(result.get("translation", "")).strip()
                note = str(result.get("note", "")).strip()
                if not core.tokens_match(rep.corrected_english, translation):
                    for item in members:
                        append_qa(
                            qa, lang, item.row, item.key, item.original_english,
                            item.corrected_english, translation,
                            f"Placeholder mismatch: source={core.tokens(item.corrected_english)} "
                            f"target={core.tokens(translation)}",
                        )
                    continue

                advisories = []
                if glossary_check:
                    violations = core.glossary_violations(rep.corrected_english, translation, terms)
                    if violations:
                        advisories.append("Glossary not applied: " + ", ".join(violations))
                if length_check:
                    ratio = core.length_ratio_flag(rep.corrected_english, translation)
                    if ratio:
                        advisories.append(ratio)
                advisory_note = "; ".join(advisories)

                cache["translations"][lang][group_key] = {"translation": translation, "note": note}

                for item in members:
                    existing = ws.cell(item.row, target_col).value
                    ws.cell(item.row, target_col).value = translation
                    combined = "; ".join(p for p in (note, advisory_note) if p)
                    if existing not in (None, "") and existing != translation:
                        append_qa(
                            qa, lang, item.row, item.key, item.original_english,
                            item.corrected_english, translation,
                            "Existing translation revised because corrected English changed"
                            + (f"; {combined}" if combined else ""),
                        )
                    elif combined:
                        append_qa(
                            qa, lang, item.row, item.key, item.original_english,
                            item.corrected_english, translation, combined,
                        )

            checkpoint()
            print(f"  Saved batch {i + 1}/{total} ({len(batch)} items)", flush=True)

    qa.freeze_panes = "A2"
    qa.auto_filter.ref = qa.dimensions
    checkpoint()

    print("\nCost optimization summary", flush=True)
    print(f"  English cache hits: {stats['english_cache_hits']}", flush=True)
    print(f"  English API items: {stats['english_api_items']}", flush=True)
    for lang in ("ja", "tw", "kr"):
        print(
            f"  {lang}: cache hits={stats['translation_cache_hits'][lang]}, "
            f"API items={stats['translation_api_items'][lang]}",
            flush=True,
        )
    print(f"  Duplicate rows avoided: {stats['deduplicated_rows']}", flush=True)
    print(f"Complete: {output_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
